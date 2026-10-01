"""
Build Word (.docx) and PDF user guide from Markdown source.
Usage: python scripts/build-user-guide.py
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "user-guide"
SOURCE = DOCS / "RENT-A-CAR-MANAGEMENT-SYSTEM-USER-GUIDE.md"
DOCX_OUT = DOCS / "RENT-A-CAR-MANAGEMENT-SYSTEM-USER-GUIDE.docx"
PDF_OUT = DOCS / "RENT-A-CAR-MANAGEMENT-SYSTEM-USER-GUIDE.pdf"


def build_docx() -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.shared import Inches, Pt
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.shared import Inches, Pt

    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    for style_name, size in [("Normal", 11), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        if style_name in styles:
            s = styles[style_name]
            s.font.name = "Calibri"
            s.font.size = Pt(size)
            if style_name == "Normal":
                s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                s.paragraph_format.line_spacing = 1.15
                s.paragraph_format.space_after = Pt(8)

    def add_paragraph(line: str, style: str = "Normal") -> None:
        p = doc.add_paragraph(style=style)
        parts = re.split(r"(\*\*[^*]+\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part:
                p.add_run(part)

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Heading 1")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("|") and "---" not in line:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and cells[0] != "Section":
                row = doc.add_table(rows=1, cols=len(cells)).rows[0]
                for i, cell in enumerate(cells):
                    row.cells[i].text = cell
        elif line.startswith("---"):
            doc.add_paragraph("_" * 40)
        else:
            add_paragraph(line)

    doc.save(DOCX_OUT)
    print(f"Created: {DOCX_OUT}")


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def build_pdf() -> None:
    result = subprocess.run(
        ["npx", "--yes", "md-to-pdf", str(SOURCE)],
        cwd=DOCS,
        shell=True,
        capture_output=True,
        text=True,
    )
    if result.returncode == 0 and PDF_OUT.exists():
        print(f"Created: {PDF_OUT}")
        return
    print("md-to-pdf failed:", result.stderr or result.stdout)
    sys.exit(1)


if __name__ == "__main__":
    if not SOURCE.exists():
        print(f"Missing source: {SOURCE}")
        sys.exit(1)
    build_docx()
    build_pdf()
